import io
import os
import csv
import logging
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload

logger = logging.getLogger(__name__)

DRIVE_SCOPES = ["https://www.googleapis.com/auth/drive.readonly"]

GOOGLE_MIME_EXPORT_MAP = {
    "application/vnd.google-apps.document": ("text/plain", ".txt"),
    "application/vnd.google-apps.spreadsheet": ("text/csv", ".csv"),
    "application/vnd.google-apps.presentation": ("text/plain", ".txt"),
}

BINARY_MIME_PARSERS = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
    "application/vnd.ms-powerpoint": "pptx",
}

TEXT_MIME_PREFIXES = ("text/",)


class DriveReader:
    def __init__(self, folder_ids: list):
        self.folder_ids = folder_ids
        self._service = None

    def _get_service(self):
        if self._service is None:
            json_path = os.environ.get("GOOGLE_SERVICE_ACCOUNT_JSON")
            if not json_path or not os.path.exists(json_path):
                raise FileNotFoundError(
                    f"Service account JSON not found at: {json_path}"
                )
            creds = service_account.Credentials.from_service_account_file(
                json_path, scopes=DRIVE_SCOPES
            )
            self._service = build("drive", "v3", credentials=creds)
        return self._service

    def _list_files_in_folder(self, folder_id: str) -> list:
        service = self._get_service()
        files = []
        page_token = None
        while True:
            response = (
                service.files()
                .list(
                    q=f"'{folder_id}' in parents and trashed=false",
                    spaces="drive",
                    fields="nextPageToken, files(id, name, mimeType)",
                    pageToken=page_token,
                )
                .execute()
            )
            for item in response.get("files", []):
                if item["mimeType"] == "application/vnd.google-apps.folder":
                    files.extend(self._list_files_in_folder(item["id"]))
                else:
                    files.append(item)
            page_token = response.get("nextPageToken")
            if not page_token:
                break
        return files

    def _export_google_file(self, file_id: str, mime_type: str) -> bytes:
        service = self._get_service()
        request = service.files().export_media(fileId=file_id, mimeType=mime_type)
        buffer = io.BytesIO()
        downloader = MediaIoBaseDownload(buffer, request)
        done = False
        while not done:
            _, done = downloader.next_chunk()
        return buffer.getvalue()

    def _download_file(self, file_id: str) -> bytes:
        service = self._get_service()
        request = service.files().get_media(fileId=file_id)
        buffer = io.BytesIO()
        downloader = MediaIoBaseDownload(buffer, request)
        done = False
        while not done:
            _, done = downloader.next_chunk()
        return buffer.getvalue()

    def _parse_docx(self, data: bytes) -> str:
        from docx import Document
        doc = Document(io.BytesIO(data))
        return "\n".join(para.text for para in doc.paragraphs)

    def _parse_xlsx(self, data: bytes) -> str:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts = []
        for sheet in wb.worksheets:
            parts.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                line = ", ".join(str(c) if c is not None else "" for c in row)
                if line.strip(", "):
                    parts.append(line)
        return "\n".join(parts)

    def _parse_pptx(self, data: bytes) -> str:
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        parts = []
        for i, slide in enumerate(prs.slides, 1):
            parts.append(f"Slide {i}:")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    parts.append(shape.text.strip())
        return "\n".join(parts)

    def _extract_content(self, file_info: dict) -> str | None:
        file_id = file_info["id"]
        mime_type = file_info["mimeType"]

        if mime_type in GOOGLE_MIME_EXPORT_MAP:
            export_mime, _ = GOOGLE_MIME_EXPORT_MAP[mime_type]
            data = self._export_google_file(file_id, export_mime)
            return data.decode("utf-8", errors="replace")

        if mime_type in BINARY_MIME_PARSERS:
            parser_type = BINARY_MIME_PARSERS[mime_type]
            data = self._download_file(file_id)
            if parser_type == "docx":
                return self._parse_docx(data)
            elif parser_type == "xlsx":
                return self._parse_xlsx(data)
            elif parser_type == "pptx":
                return self._parse_pptx(data)

        if any(mime_type.startswith(prefix) for prefix in TEXT_MIME_PREFIXES):
            data = self._download_file(file_id)
            return data.decode("utf-8", errors="replace")

        return None

    def load_all_content(self) -> list:
        results = []
        for folder_id in self.folder_ids:
            try:
                files = self._list_files_in_folder(folder_id)
            except Exception as e:
                logger.error(f"Failed to list folder {folder_id}: {e}")
                continue

            for file_info in files:
                try:
                    content = self._extract_content(file_info)
                    if content is None:
                        logger.info(f"Skipping unsupported file type: {file_info['name']} ({file_info['mimeType']})")
                        continue
                    results.append({
                        "filename": file_info["name"],
                        "content": content,
                        "mime_type": file_info["mimeType"],
                    })
                    logger.info(f"Loaded: {file_info['name']} ({len(content)} chars)")
                except Exception as e:
                    logger.error(f"Failed to load {file_info['name']}: {e}")

        return results
